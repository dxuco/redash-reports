# Builds streamers-data.json from the query-1732 month caches.
#
#   python build_streamers.py
#   python make_streamers_html.py
#
# THE COHORT
# ----------
# Every player whose FIRST DEPOSIT landed in 2026 through a streamer:
# a row with `ftd` > 0 and `aff_type` == "Streamer". The streamer credited is
# the `aff_username` on that FTD row — attribution is fixed at acquisition, so
# a player never moves between streamers later in the year.
#
# `ftd` in the cache is a DOLLAR AMOUNT, not a flag. Count distinct player_id
# with ftd > 0; summing the field gives the first-deposit value, never a count.
#
# Deposits, bets, GGR, NGR and adjusted GGR are then summed over every 2026 row
# belonging to those players — so the money columns are "what this cohort did
# during 2026", not "what happened in the month they were acquired".
#
# THE WHALE
# ---------
# The carved-out players (see ../data-exclusions.json) are held out of every
# other report by default.
# He is NOT in this cohort — he first deposited years before 2026 and not via a
# streamer — so there is no toggle here. The assertion below keeps that true: if
# he ever appears, the build fails rather than quietly inflating deposits 3-4x.

import json, os, glob
from collections import defaultdict
from datetime import date as _date

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
CACHE = os.path.join(ROOT, "ftd-report", "cache")

YEAR = "2026"
# The players every report carves out by default, shared so a tenth report
# cannot quietly disagree about who is in the list. See ../carveout.py.
import sys
sys.path.insert(0, ROOT)
from carveout import CARVED as WHALE_SET, CARVE_NAMES
UNTAGGED = "(untagged streamer)"

# Cohort-quality variants. The page switches between them; each is a different
# COHORT, not a different filter on one cohort, so every figure — deposits,
# adjusted GGR, the streamer count itself — is recomputed for each.
#   all -> every 2026 streamer FTD
#   q   -> Qualified and Super Qualified only
#   sq  -> Super Qualified only
VARIANTS = {
    "all": {"Non Qualified", "Qualified", "Super Qualified"},
    "q":   {"Qualified", "Super Qualified"},
    "sq":  {"Super Qualified"},
}
TYPES = ["Non Qualified", "Qualified", "Super Qualified"]


def num(v):
    if v in (None, "", "None"):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# --- load -------------------------------------------------------------------
paths = sorted(p for p in glob.glob(os.path.join(CACHE, "*.json"))
               if os.path.basename(p).startswith(YEAR))
assert paths, "no %s month caches found in %s" % (YEAR, CACHE)
MONTHS = [os.path.basename(p)[:-5] for p in paths]
rows_by_month = {}
last_day = ""
for p in paths:
    m = os.path.basename(p)[:-5]
    rows = json.load(open(p, encoding="utf-8"))
    rows_by_month[m] = rows
    for r in rows:
        d = (r.get("transaction_date") or "")[:10]
        if d > last_day:
            last_day = d

# --- the cohort -------------------------------------------------------------
# pid -> [aff, ftd_type, ftd month, ftd value, country]
cohort = {}
cohort_fdd = {}          # pid -> first_deposit_date, YYYY-MM-DD
dupes = 0
untagged = 0
for m in MONTHS:
    for r in rows_by_month[m]:
        if r.get("aff_type") != "Streamer":
            continue
        v = num(r.get("ftd"))
        if v <= 0:
            continue
        pid = r.get("player_id")
        if pid in cohort:
            # An FTD is once per player by definition; a second one is a data
            # fault, not a second acquisition. Keep the first, count the rest.
            dupes += 1
            continue
        aff = r.get("aff_username") or UNTAGGED
        if aff == UNTAGGED:
            untagged += 1
        cohort[pid] = [aff, r.get("ftd_type") or "Non Qualified", m, v,
                       r.get("player_country") or "Unknown"]
        # Day 0 for retention. `first_deposit_date` is carried on every row the
        # player has and never disagrees with itself, so it is taken at face
        # value rather than re-derived — the same call retention/ makes.
        fdd = (r.get("first_deposit_date") or "")[:10]
        if fdd:
            cohort_fdd[pid] = fdd

# This report has no carve-out toggle, and that is only defensible while none
# of the carved-out players is in the cohort. Naming whoever turns up is the
# point: the fix is to add a toggle, and you cannot add one for a player the
# message does not identify.
_in_cohort = sorted(WHALE_SET & set(cohort))
assert not _in_cohort, (
    "carved-out player(s) %s are in the %s streamer FTD cohort — this report "
    "has no carve-out toggle because none of them was ever in it. Add one "
    "before shipping." % (", ".join("%s (%s)" % (i, CARVE_NAMES.get(i, "?"))
                                    for i in _in_cohort), YEAR))

# --- money, per streamer and per month --------------------------------------
# agg[variant][aff] = dict of totals;  ser[variant][aff][month] = dict
def blank():
    return {"ftd": 0, "nq": 0, "q": 0, "sq": 0, "ftdv": 0.0,
            "dep": 0.0, "depn": 0, "wdr": 0.0, "bet": 0.0,
            "ggr": 0.0, "ngr": 0.0, "adj": 0.0, "bonus": 0.0,
            "ggrPos": 0.0, "ngrNeg": 0.0, "nWin": 0, "nLose": 0,
            "rK": [0, 0], "rN": [0, 0],
            "players": set(), "countries": defaultdict(int)}


agg = {v: defaultdict(blank) for v in VARIANTS}
# ser[variant][(aff, month)] = [ftd, sq, dep, adj, ngr]
ser = {v: defaultdict(lambda: [0, 0, 0.0, 0.0, 0.0]) for v in VARIANTS}
# pid -> [ggr, ngr] over the whole year, and the set of days they deposited on
player_money = defaultdict(lambda: [0.0, 0.0])
dep_days = defaultdict(set)

TYPE_KEY = {"Non Qualified": "nq", "Qualified": "q", "Super Qualified": "sq"}

for pid, (aff, ftype, fm, fv, ctry) in cohort.items():
    for var, keep in VARIANTS.items():
        if ftype not in keep:
            continue
        a = agg[var][aff]
        a["ftd"] += 1
        a[TYPE_KEY[ftype]] += 1
        a["ftdv"] += fv
        a["players"].add(pid)
        a["countries"][ctry] += 1
        s = ser[var][(aff, fm)]
        s[0] += 1
        if ftype == "Super Qualified":
            s[1] += 1

money_rows = 0
for m in MONTHS:
    for r in rows_by_month[m]:
        pid = r.get("player_id")
        c = cohort.get(pid)
        if not c:
            continue
        aff, ftype = c[0], c[1]
        dep, adj, ngr = num(r.get("deposit")), num(r.get("adjusted_ggr")), num(r.get("ngr"))
        ggr, bet = num(r.get("ggr")), num(r.get("bet"))
        bon, wdr = num(r.get("bonus_cost")), num(r.get("withdraw"))
        dn = int(num(r.get("deposit_count")))
        money_rows += 1
        # Per player, for the winners/losers split below. A player is a winner or
        # a loser on their WHOLE-year result, not day by day.
        pm = player_money[pid]
        pm[0] += ggr; pm[1] += ngr
        if dep:
            dep_days[pid].add((r.get("transaction_date") or "")[:10])
        for var, keep in VARIANTS.items():
            if ftype not in keep:
                continue
            a = agg[var][aff]
            a["dep"] += dep; a["depn"] += dn; a["wdr"] += wdr
            a["bet"] += bet; a["ggr"] += ggr; a["ngr"] += ngr
            a["adj"] += adj; a["bonus"] += bon
            s = ser[var][(aff, m)]
            s[2] += dep; s[3] += adj; s[4] += ngr

# --- the winners / losers split ---------------------------------------------
# "Positive GGR" is what the streamer earned from the players who finished 2026
# GGR-positive; "negative NGR" is what the players who finished NGR-negative
# took back. Split at the PLAYER, not the row: a profitable player still has
# losing days, and summing signed rows instead would inflate both sides while
# telling you nothing about who the streamer actually sent.
for pid, (ggr, ngr) in player_money.items():
    aff, ftype = cohort[pid][0], cohort[pid][1]
    for var, keep in VARIANTS.items():
        if ftype not in keep:
            continue
        a = agg[var][aff]
        if ggr > 0:
            a["ggrPos"] += ggr
            a["nWin"] += 1
        if ngr < 0:
            a["ngrNeg"] += ngr
            a["nLose"] += 1

# --- deposit retention ------------------------------------------------------
# The definition is taken VERBATIM from retention/build_retention.py, because
# two pages on this site quoting different retention numbers is a worse bug than
# either being slightly wrong:
#
#   day 0    = the player's first_deposit_date
#   retained = a deposit on some day 1..N after it (day-0 top-ups are not a
#              return, and betting is not retention — the question is whether
#              they fund the account again)
#   eligible = only players observed for at least N days. A player who first
#              deposited four days before the cache ends cannot fail day 7, so
#              counting them in the denominator would drag every recent
#              streamer's rate down for no reason.
CUTOFF = _date(int(last_day[0:4]), int(last_day[5:7]), int(last_day[8:10]))
MILESTONES = [7, 30]
# retention/build_retention.py uses 30 as the floor below which a rate is noise.
# Per streamer the cohorts are an order of magnitude smaller — 30 would blank
# most of the table — so the floor is 10 here, and the denominator is on every
# cell as a tooltip so a thin rate can be recognised as one.
MIN_ELIGIBLE = 10

for pid, (aff, ftype, fm, fv, ctry) in cohort.items():
    fd = cohort_fdd.get(pid)
    if not fd:
        continue
    f = _date(int(fd[0:4]), int(fd[5:7]), int(fd[8:10]))
    obs = (CUTOFF - f).days
    if obs < 0:
        # First deposit dated after the last day any cache covers: their day 0
        # has not been loaded, so every day-N figure for them measures nothing.
        continue
    lags = [(_date(int(x[0:4]), int(x[5:7]), int(x[8:10])) - f).days
            for x in dep_days.get(pid, ())]
    first_back = min([g for g in lags if g >= 1], default=None)
    for var, keep in VARIANTS.items():
        if ftype not in keep:
            continue
        a = agg[var][aff]
        for j, n in enumerate(MILESTONES):
            if obs < n:
                continue
            a["rN"][j] += 1
            if first_back is not None and first_back <= n:
                a["rK"][j] += 1

# --- investment, from the CPA cost sheet ------------------------------------
# What we PAID each streamer. It is not in query 1732 and never will be — it
# lives in the CPA deal sheet the Acquisition report already reads.
#
# The file is whatever step 4 (build_acquisition.py) last downloaded from Drive,
# with the same COST_DIR fallback. Reading its copy rather than downloading
# again means this page cannot disagree with /acquisition-2026 about what a
# streamer cost, and costs one fewer Drive round-trip. Step 4 runs before step 9
# in UPDATE-EVERYTHING, so the copy is same-morning fresh.
#
# The match is by `Affiliate Name` against `aff_username`, case-insensitive —
# copied from acquisition-report/gen_html.py, which is the only other place that
# joins these two datasets. Two pages disagreeing about which deal belongs to
# which streamer would be worse than either being slightly wrong.
ACQ = os.path.join(ROOT, "acquisition-report")


def read_config(path, key, default=""):
    if not os.path.exists(path):
        return default
    for raw in open(path, encoding="utf-8"):
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        if k.strip() == key:
            return v.strip().strip('"').strip("'")
    return default


def find_cost_file():
    drive_copy = os.path.join(ACQ, "cost_from_drive.xlsx")
    if os.path.exists(drive_copy):
        return drive_copy
    cost_dir = read_config(os.path.join(ACQ, "config.env"), "COST_DIR")
    if cost_dir and os.path.isdir(cost_dir):
        xls = sorted(glob.glob(os.path.join(cost_dir, "*.xlsx")),
                     key=os.path.getmtime, reverse=True)
        if xls:
            return xls[0]
    return None


invest = defaultdict(float)          # aff_username.lower() -> cost paid
deals = defaultdict(int)
cost_file = find_cost_file()
cost_note = ""
if cost_file is None:
    cost_note = "no cost file found — investment is blank on the page"
else:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(cost_file, read_only=True, data_only=True)
        # The 2026 deal sheet was called 'Data' until Aug 2026, then renamed
        # 'Data2026' when a 'Data2025' sheet was added. Accept either.
        sheet = next((n for n in ("Data2026", "Data") if n in wb.sheetnames), None)
        if sheet is None:
            cost_note = "cost file has no Data2026/Data sheet: %s" % ", ".join(wb.sheetnames)
        else:
            crows = list(wb[sheet].iter_rows(min_row=1, values_only=True))
            ch = [(h or "").strip() for h in crows[0]]
            ci = {h: i for i, h in enumerate(ch) if h}

            def cg(r, k):
                i = ci.get(k)
                return "" if i is None or i >= len(r) or r[i] is None else str(r[i]).strip()

            for r in crows[1:]:
                if not any(c is not None and str(c).strip() != "" for c in r[:11]):
                    continue
                nm = cg(r, "Affiliate Name")
                if not nm:
                    continue
                invest[nm.lower()] += num(cg(r, "Cost"))
                deals[nm.lower()] += 1
            cost_note = "%s, %d affiliates" % (os.path.basename(cost_file), len(invest))
    except ImportError:
        cost_note = "openpyxl not installed (pip install openpyxl) — investment is blank"
    except Exception as e:                       # a bad sheet must not take the page down
        cost_note = "could not read %s: %s" % (os.path.basename(cost_file), e)

# --- shape ------------------------------------------------------------------
R2 = lambda x: round(x, 2)


def top_country(counts):
    if not counts:
        return ["Unknown", 0]
    k = max(counts.items(), key=lambda kv: (kv[1], kv[0]))
    return [k[0], k[1]]


out_variants = {}
for var in VARIANTS:
    streamers = []
    for aff, a in agg[var].items():
        if a["ftd"] == 0:
            continue
        tc = top_country(a["countries"])
        streamers.append({
            "a": aff, "ftd": a["ftd"], "nq": a["nq"], "q": a["q"], "sq": a["sq"],
            "ftdv": R2(a["ftdv"]), "dep": R2(a["dep"]), "depn": a["depn"],
            "wdr": R2(a["wdr"]), "bet": R2(a["bet"]), "ggr": R2(a["ggr"]),
            "ngr": R2(a["ngr"]), "adj": R2(a["adj"]), "bonus": R2(a["bonus"]),
            "ggrPos": R2(a["ggrPos"]), "ngrNeg": R2(a["ngrNeg"]),
            "nWin": a["nWin"], "nLose": a["nLose"],
            # retained / eligible at each milestone, kept as the two counts so
            # the page can sum them into a totals row. A rate cannot be averaged
            # across streamers; the counts can.
            "r7": a["rK"][0], "r7n": a["rN"][0],
            "r30": a["rK"][1], "r30n": a["rN"][1],
            # `inv` is null, not 0, when the streamer has no row in the cost
            # sheet at all. $0 spent and "we have no record" are different
            # facts, and averaging the second in as a zero would understate
            # every cost-per-FTD figure on the page.
            "inv": R2(invest[aff.lower()]) if aff.lower() in invest else None,
            "deals": deals.get(aff.lower(), 0),
            "ctry": tc[0], "ctryN": tc[1],
            "m": sorted({m for (x, m) in ser[var] if x == aff}),
        })
    # deposits descending is the order a reader wants first; the page re-sorts.
    streamers.sort(key=lambda s: -s["dep"])

    # Monthly series. FTD counts sit on the ACQUISITION month; money sits on the
    # month it happened. Two different months on one chart would be a lie, so
    # they are emitted separately and captioned separately on the page.
    idx = {m: i for i, m in enumerate(MONTHS)}
    z = lambda: [0] * len(MONTHS)
    ftdM = {t: z() for t in TYPES}
    depM, adjM, ngrM = z(), z(), z()
    for pid, (aff, ftype, fm, fv, ctry) in cohort.items():
        if ftype not in VARIANTS[var]:
            continue
        ftdM[ftype][idx[fm]] += 1
    for (aff, m), s in ser[var].items():
        i = idx[m]
        depM[i] += s[2]; adjM[i] += s[3]; ngrM[i] += s[4]

    pos = [s for s in streamers if s["adj"] > 0]
    neg = [s for s in streamers if s["adj"] < 0]
    zero = [s for s in streamers if s["adj"] == 0]

    out_variants[var] = {
        "streamers": streamers,
        "series": {
            "ftd": {t: ftdM[t] for t in TYPES},
            "dep": [R2(x) for x in depM],
            "adj": [R2(x) for x in adjM],
            "ngr": [R2(x) for x in ngrM],
        },
        # No per-month positive/negative split is emitted. The page splits by
        # the streamer's WHOLE-YEAR result and derives the bands from `cells`,
        # and a second, month-by-month split sitting here unused would read like
        # the one the chart draws — it is not, and the two disagree by design.
        "totals": {
            "streamers": len(streamers),
            "ftd": sum(s["ftd"] for s in streamers),
            "sq": sum(s["sq"] for s in streamers),
            "q": sum(s["q"] for s in streamers),
            "nq": sum(s["nq"] for s in streamers),
            "ftdv": R2(sum(s["ftdv"] for s in streamers)),
            "dep": R2(sum(s["dep"] for s in streamers)),
            "depn": sum(s["depn"] for s in streamers),
            "adj": R2(sum(s["adj"] for s in streamers)),
            "ngr": R2(sum(s["ngr"] for s in streamers)),
            "ggr": R2(sum(s["ggr"] for s in streamers)),
            "bet": R2(sum(s["bet"] for s in streamers)),
            "bonus": R2(sum(s["bonus"] for s in streamers)),
            "ggrPos": R2(sum(s["ggrPos"] for s in streamers)),
            "ngrNeg": R2(sum(s["ngrNeg"] for s in streamers)),
            "nWin": sum(s["nWin"] for s in streamers),
            "nLose": sum(s["nLose"] for s in streamers),
            "r7": sum(s["r7"] for s in streamers), "r7n": sum(s["r7n"] for s in streamers),
            "r30": sum(s["r30"] for s in streamers), "r30n": sum(s["r30n"] for s in streamers),
            "inv": R2(sum(s["inv"] or 0 for s in streamers)),
            "nInv": sum(1 for s in streamers if s["inv"] is not None),
            # FTDs belonging to streamers we have a cost row for. Cost per FTD
            # must divide by these, not by every FTD, or the 17 streamers with
            # no deal on file quietly discount the whole programme's CPA.
            "invFtd": sum(s["ftd"] for s in streamers if s["inv"] is not None),
            "nPos": len(pos), "nNeg": len(neg), "nZero": len(zero),
            "adjPos": R2(sum(s["adj"] for s in pos)),
            "adjNeg": R2(sum(s["adj"] for s in neg)),
        },
        # [aff, month, ftd, sq, dep, adj, ngr] — the per-streamer drill row
        "cells": [[aff, m, s[0], s[1], R2(s[2]), R2(s[3]), R2(s[4])]
                  for (aff, m), s in sorted(ser[var].items())],
    }

data = {
    "generated": last_day,
    "year": YEAR,
    "months": MONTHS,
    "types": TYPES,
    "untagged": UNTAGGED,
    "cohortPlayers": len(cohort),
    "milestones": MILESTONES,
    "minEligible": MIN_ELIGIBLE,
    "costFile": os.path.basename(cost_file) if cost_file else None,
    "costNote": cost_note,
    "cutoff": last_day,
    "variants": out_variants,
}

dest = os.path.join(HERE, "streamers-data.json")
with open(dest, "w", encoding="utf-8") as f:
    json.dump(data, f, separators=(",", ":"))

# --- diagnostics ------------------------------------------------------------
print("months %s .. %s (data to %s)" % (MONTHS[0], MONTHS[-1], last_day))
print("cohort: %d players, %d repeat FTD rows skipped, %d untagged"
      % (len(cohort), dupes, untagged))
print("carved-out players absent from cohort: OK (%s)" % ", ".join(sorted(WHALE_SET)))
print("money rows attributed: %d" % money_rows)
print("cost sheet: %s" % (cost_note or "none"))
print()
for var in ("all", "q", "sq"):
    t = out_variants[var]["totals"]
    print("%-3s  streamers %3d  FTD %5d (SQ %4d)  ftd$ %11s  dep$ %12s  adjGGR$ %11s"
          % (var, t["streamers"], t["ftd"], t["sq"],
             format(round(t["ftdv"]), ","), format(round(t["dep"]), ","),
             format(round(t["adj"]), ",")))
    print("     positive %3d (+%s)   negative %3d (%s)   flat %d"
          % (t["nPos"], format(round(t["adjPos"]), ","),
             t["nNeg"], format(round(t["adjNeg"]), ","), t["nZero"]))
    print("     GGR from the %4d winning players +%-11s   NGR to the %4d losing players %s"
          % (t["nWin"], format(round(t["ggrPos"]), ","),
             t["nLose"], format(round(t["ngrNeg"]), ",")))
    print("     invested %-11s across %3d of %3d streamers   cost/FTD %s on their %s FTDs"
          % ("$" + format(round(t["inv"]), ","), t["nInv"], t["streamers"],
             "$%.0f" % (t["inv"] / t["invFtd"]) if t["invFtd"] else "n/a",
             format(t["invFtd"], ",")))
    print("     D7 %s of %s eligible (%.0f%%)   D30 %s of %s (%.0f%%)"
          % (format(t["r7"], ","), format(t["r7n"], ","),
             100.0 * t["r7"] / t["r7n"] if t["r7n"] else 0,
             format(t["r30"], ","), format(t["r30n"], ","),
             100.0 * t["r30"] / t["r30n"] if t["r30n"] else 0))

# Reconcile: the per-streamer rows and the monthly series are two independent
# roll-ups of the same cohort, so they must agree inside the rounding budget.
for var in VARIANTS:
    v = out_variants[var]
    n = len(v["streamers"])
    budget = (n + len(MONTHS)) * 0.005
    for field, series in (("dep", v["series"]["dep"]), ("adj", v["series"]["adj"]),
                          ("ngr", v["series"]["ngr"])):
        a = sum(s[field] for s in v["streamers"])
        b = sum(series)
        assert abs(a - b) <= budget, \
            "%s/%s: streamer rows %.2f vs monthly series %.2f (budget %.2f)" \
            % (var, field, a, b, budget)
    ftd_rows = sum(s["ftd"] for s in v["streamers"])
    ftd_ser = sum(sum(x) for x in v["series"]["ftd"].values())
    assert ftd_rows == ftd_ser, "%s: FTD %d vs %d" % (var, ftd_rows, ftd_ser)
    assert v["totals"]["nPos"] + v["totals"]["nNeg"] + v["totals"]["nZero"] == n
print("\nreconciled: streamer rows == monthly series, for every variant")
print("wrote streamers-data.json (%.0f KB)" % (os.path.getsize(dest) / 1024.0))
